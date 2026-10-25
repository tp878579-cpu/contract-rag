# src/ingestion/loader.py
"""Load PDF / DOCX / TXT files and return text page by page."""

from dataclasses import dataclass
from pathlib import Path

import pymupdf  # PyMuPDF
from docx import Document

from src.utils.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


@dataclass
class PageContent:
    """Text of one page, plus where it came from (used later for citations)."""
    text: str
    page_number: int      # starts at 1
    source: str           # file name, e.g. "nda_acme.pdf"


def _load_pdf(path: Path) -> list[PageContent]:
    pages = []
    with pymupdf.open(path) as pdf:
        for i, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if text:  # skip blank pages
                pages.append(PageContent(text=text, page_number=i, source=path.name))
    if not pages:
        logger.warning("%s has no extractable text (scanned PDF? needs OCR)", path.name)
    return pages


def _load_docx(path: Path) -> list[PageContent]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Contracts often keep important terms inside tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    # DOCX has no fixed pages, so the whole file counts as page 1
    return [PageContent(text=text, page_number=1, source=path.name)] if text else []


def _load_txt(path: Path) -> list[PageContent]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [PageContent(text=text, page_number=1, source=path.name)] if text else []


def load_document(file_path: str | Path) -> list[PageContent]:
    """Load one file and return a list of PageContent."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Use: {sorted(SUPPORTED_EXTENSIONS)}")

    loaders = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_txt}
    pages = loaders[ext](path)

    logger.info("Loaded %s: %d page(s)", path.name, len(pages))
    return pages


def load_folder(folder: str | Path) -> list[PageContent]:
    """Load every supported file in a folder."""
    all_pages = []
    for file in sorted(Path(folder).iterdir()):
        if file.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                all_pages.extend(load_document(file))
            except Exception as e:
                logger.error("Failed to load %s: %s", file.name, e)
    return all_pages